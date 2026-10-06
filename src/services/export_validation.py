"""Validate exported DOCX/PDF files: they open, contain every approved line, and look ATS-safe.

PDF checks use the extracted text and character geometry (pdfplumber), which is
what a parser sees. DOCX page count can't be measured without a word processor,
so it is reported from the PDF, which carries the same content.
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from src.schemas.export import CheckStatus, ExportCheck
from src.services.render.blocks import Block
from src.services.render.style import MIN_FONT_PT, PdfFont, unrenderable_chars

_DASHES = dict.fromkeys(map(ord, "‐‑‒–—−"), "-")
_QUOTES = {ord("‘"): "'", ord("’"): "'", ord("“"): '"', ord("”"): '"'}
MAX_LISTED = 5


def normalize(text: str) -> str:
    """Compare text the way a parser would: ignore case, dash/quote styles, bullets and line wrapping."""
    text = unicodedata.normalize("NFKC", text).translate(_DASHES).translate(_QUOTES)
    text = text.replace("­", "").replace("•", " ")
    return re.sub(r"\s+", " ", text).strip().casefold()


def _missing(blocks: list[Block], extracted: str) -> list[str]:
    haystack = normalize(extracted)
    return [b.text for b in blocks if normalize(b.text) not in haystack]


def _listed(items: list[str]) -> str:
    shown = "; ".join(f"“{i[:60]}”" for i in items[:MAX_LISTED])
    return shown + (f"; and {len(items) - MAX_LISTED} more" if len(items) > MAX_LISTED else "")


def _check(file: str, name: str, ok: bool, detail_ok: str, detail_bad: str,
           bad_status: CheckStatus = CheckStatus.FAIL) -> ExportCheck:
    return ExportCheck(file=file, name=name, status=CheckStatus.PASS if ok else bad_status,
                       detail=detail_ok if ok else detail_bad)


def validate_pdf(path: Path, blocks: list[Block], target_pages: int, font: PdfFont) -> tuple[list[ExportCheck], int | None]:
    import pdfplumber

    checks: list[ExportCheck] = []
    try:
        pdf = pdfplumber.open(path)
    except Exception as exc:  # any failure to open is a failed export
        return [ExportCheck(file="pdf", name="Opens", status=CheckStatus.FAIL,
                            detail=f"The PDF could not be opened ({type(exc).__name__}).")], None
    with pdf:
        pages = len(pdf.pages)
        texts = [(p.extract_text() or "") for p in pdf.pages]
        chars = [(p, ch) for p in pdf.pages for ch in p.chars]
        checks.append(ExportCheck(file="pdf", name="Opens", status=CheckStatus.PASS, detail=f"{pages} page(s)."))
        blank = [i + 1 for i, t in enumerate(texts) if not t.strip()]
        checks.append(_check("pdf", "No blank pages", not blank, "Every page has text.",
                             f"Blank page(s): {', '.join(map(str, blank))}."))
        missing = _missing(blocks, "\n".join(texts))
        checks.append(_check("pdf", "All approved content present", not missing,
                             f"All {len(blocks)} approved lines found in the extracted text.",
                             f"Missing or unreadable in the PDF text: {_listed(missing)}."))
        off = sum(1 for p, ch in chars if ch["x0"] < 0 or ch["x1"] > p.width or ch["top"] < 0 or ch["bottom"] > p.height)
        checks.append(_check("pdf", "Text stays on the page", off == 0, "No text runs past the page edges.",
                             f"{off} character(s) run off the page."))
        smallest = min((ch["size"] for _, ch in chars), default=0)
        checks.append(_check("pdf", "Readable font size", smallest >= MIN_FONT_PT - 0.01,
                             f"Smallest text is {smallest:.1f} pt.",
                             f"Some text is {smallest:.1f} pt (minimum {MIN_FONT_PT:g} pt)."))
        cid = sum(t.count("(cid:") for t in texts)
        checks.append(_check("pdf", "Text extracts cleanly", cid == 0, "No unmapped glyphs in the extracted text.",
                             f"{cid} glyph(s) extract as unreadable codes; parsers may see junk."))
    bad = sorted(unrenderable_chars(" ".join(b.text for b in blocks), font))
    checks.append(_check("pdf", "Characters supported by font", not bad,
                         f"Font: {font.source}.",
                         f"The fallback font can't draw: {' '.join(bad)}. Set RESUME_OPTIMIZER_PDF_FONT to a "
                         "TrueType font that has them."))
    checks.append(_check("pdf", "Fits target length", pages <= target_pages,
                         f"{pages} page(s) for a {target_pages}-page target.",
                         f"{pages} pages for a {target_pages}-page target. Text is never shrunk to fit; remove "
                         "bullets in Edit, or generate with the 2-page target.", CheckStatus.WARN))
    return checks, pages


def validate_docx(path: Path, blocks: list[Block]) -> list[ExportCheck]:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(str(path))
    except (PackageNotFoundError, ValueError, KeyError, OSError) as exc:
        return [ExportCheck(file="docx", name="Opens", status=CheckStatus.FAIL,
                            detail=f"The DOCX could not be opened ({type(exc).__name__}).")]
    checks = [ExportCheck(file="docx", name="Opens", status=CheckStatus.PASS, detail="Opens as a Word document.")]
    body = "\n".join(p.text for p in document.paragraphs)
    missing = _missing(blocks, body)
    checks.append(_check("docx", "All approved content present", not missing,
                         f"All {len(blocks)} approved lines found.", f"Missing: {_listed(missing)}."))
    checks.append(_check("docx", "No tables", not document.tables, "Single-column text, no tables.",
                         f"{len(document.tables)} table(s) found; some parsers misread tables."))
    checks.append(_check("docx", "No images or text boxes", not document.inline_shapes
                         and "txbxContent" not in document.element.xml,
                         "No images, icons or text boxes.", "Images or text boxes found.", CheckStatus.WARN))
    header_text = " ".join(p.text for s in document.sections for part in (s.header, s.footer)
                           for p in part.paragraphs).strip()
    checks.append(_check("docx", "Contact details in body", not header_text,
                         "Nothing in page headers/footers.", "Text found in a page header/footer; some parsers "
                         "skip those.", CheckStatus.WARN))
    sizes = [r.font.size.pt for p in document.paragraphs for r in p.runs if r.font.size is not None]
    sizes += [s.font.size.pt for s in document.styles if getattr(s, "font", None) is not None
              and s.font.size is not None and s.name in ("Normal", "List Bullet", "Heading 1")]
    smallest = min(sizes, default=MIN_FONT_PT)
    checks.append(_check("docx", "Readable font size", smallest >= MIN_FONT_PT - 0.01,
                         f"Smallest text is {smallest:g} pt.", f"Some text is {smallest:g} pt."))
    return checks
