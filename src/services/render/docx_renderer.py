"""DOCX renderer: single column, built-in heading and list styles, no tables/text boxes/images,
contact details in the body (not the page header)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from src.services.render.blocks import Block
from src.services.render.style import PAGE_SIZES_IN, Style

_BLACK = RGBColor(0, 0, 0)


def _set_font(style_or_run, family: str, size: float, *, bold: bool | None = None) -> None:
    font = style_or_run.font
    font.name = family
    font.size = Pt(size)
    font.color.rgb = _BLACK
    if bold is not None:
        font.bold = bold
    # Make East-Asian/complex-script text use the same family instead of a theme font.
    element = style_or_run.element
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), family)


def _bottom_border(paragraph) -> None:
    """Thin rule under a section heading (a paragraph border, not a drawing or table)."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    ppr.append(borders)


def render_docx(blocks: list[Block], path: Path, style: Style, page_size: str = "A4", *, title: str = "") -> Path:
    doc = Document()
    section = doc.sections[0]
    width, height = PAGE_SIZES_IN[page_size]
    section.page_width, section.page_height = Inches(width), Inches(height)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(style.margin_in))

    normal = doc.styles["Normal"]
    _set_font(normal, style.font_family, style.body_pt)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = style.line_spacing
    heading = doc.styles["Heading 1"]
    _set_font(heading, style.font_family, style.heading_pt, bold=True)
    heading.font.italic = False
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True
    bullet_style = doc.styles["List Bullet"]
    _set_font(bullet_style, style.font_family, style.body_pt)
    bullet_style.paragraph_format.space_after = Pt(1)

    for block in blocks:
        if block.kind == "name":
            p = doc.add_paragraph()
            _set_font(p.add_run(block.text), style.font_family, style.name_pt, bold=True)
        elif block.kind == "headline":
            p = doc.add_paragraph()
            _set_font(p.add_run(block.text), style.font_family, style.headline_pt)
        elif block.kind == "contact":
            p = doc.add_paragraph()
            _set_font(p.add_run(block.text), style.font_family, style.contact_pt)
        elif block.kind == "heading":
            p = doc.add_heading(block.text, level=1)
            _bottom_border(p)
        elif block.kind == "entry":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.keep_with_next = True
            first, sep, rest = block.text.partition(" | ")
            _set_font(p.add_run(first), style.font_family, style.body_pt, bold=True)
            if sep:
                _set_font(p.add_run(sep + rest), style.font_family, style.body_pt)
        elif block.kind == "bullet":
            doc.add_paragraph(block.text, style="List Bullet")
        else:  # text
            doc.add_paragraph(block.text)

    props = doc.core_properties
    props.title = title
    props.author = next((b.text for b in blocks if b.kind == "name"), "")
    props.comments = props.keywords = props.subject = ""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
