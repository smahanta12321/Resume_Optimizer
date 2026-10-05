"""JD import: text normalization, file extraction, URL canonicalization, deduplication.

Imported JD text is untrusted data. It is stored and displayed, never executed
or treated as instructions.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import PurePath
from typing import Optional
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from src.schemas.job import MAX_JD_CHARS, Job, SourcePlatform

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES = 5_000_000
MAX_PDF_PAGES = 30
MIN_JD_CHARS = 200  # shorter text is allowed but flagged as possibly incomplete
SUPPORTED_EXTENSIONS = (".txt", ".md", ".pdf", ".docx")

# Query parameters that only track the visitor; dropping them never changes the posting.
_TRACKING_PARAMS = {
    "trk", "trkinfo", "refid", "trackingid", "lipi", "eby", "src", "sid", "xid", "xp", "px",
    "position", "pagenum", "originalsubdomain", "ref", "referrer", "fbclid", "gclid", "msclkid",
}
_LINKEDIN_JOB_RE = re.compile(r"/jobs/view/(?:[^/]*?-)?(\d{6,})/?")
_NAUKRI_JOB_RE = re.compile(r"/job-listings-[^/?#]*?-(\d{9,})/?$")

# Near-duplicate threshold on word-shingle Jaccard similarity of descriptions.
NEAR_DUPLICATE_SIMILARITY = 0.85


class JDImportError(ValueError):
    """Raised when JD input cannot be imported. Message is user-facing."""


# ---- text -------------------------------------------------------------------


def normalize_jd_text(text: str) -> str:
    """Clean extracted/pasted JD text: NFC, no control chars, tidy whitespace, keep paragraphs."""
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(ch for ch in text if ch in "\n\t" or unicodedata.category(ch)[0] != "C")
    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def check_jd_text(text: str) -> tuple[str, list[str]]:
    """Normalize JD text and return ``(text, notes)``. Raises on empty or oversized input."""
    text = normalize_jd_text(text)
    if not text:
        raise JDImportError("The job description is empty.")
    if len(text) > MAX_JD_CHARS:
        raise JDImportError(
            f"The job description is too long ({len(text):,} characters; limit {MAX_JD_CHARS:,}). "
            "Paste only the job posting itself."
        )
    notes = []
    if len(text) < MIN_JD_CHARS:
        notes.append(
            f"The text is short ({len(text)} characters). It may be incomplete; check that the full JD was captured."
        )
    return text, notes


# ---- files ------------------------------------------------------------------


def extract_text_from_file(filename: str, data: bytes) -> str:
    """Extract plain text from an uploaded TXT/MD/PDF/DOCX file.

    The extension must be supported AND the content must match it (magic bytes),
    so a renamed file cannot reach the wrong parser.
    """
    ext = PurePath(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise JDImportError(f"Unsupported file type {ext or '(none)'!r}. Use one of: {', '.join(SUPPORTED_EXTENSIONS)}.")
    if not data:
        raise JDImportError("The file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise JDImportError(f"File is too large ({len(data):,} bytes; limit {MAX_UPLOAD_BYTES:,}).")

    if ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise JDImportError("This file does not look like a PDF.")
        return _pdf_text(data)
    if ext == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise JDImportError("This file does not look like a DOCX document.")
        return _docx_text(data)
    if b"\x00" in data[:4096]:
        raise JDImportError("This file looks binary, not plain text.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise JDImportError("Could not decode the text file. Save it as UTF-8 and try again.")


def _pdf_text(data: bytes) -> str:
    import pdfplumber
    from pdfplumber.utils.exceptions import PdfminerException

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise JDImportError(f"PDF has {len(pdf.pages)} pages; limit is {MAX_PDF_PAGES}.")
            text = "\n\n".join((page.extract_text() or "") for page in pdf.pages)
    except JDImportError:
        raise
    except (PdfminerException, ValueError, KeyError, TypeError) as exc:  # pdfplumber wraps pdfminer errors
        logger.warning("event=pdf_parse_failed error_type=%s", type(exc).__name__)
        raise JDImportError("The PDF could not be read. It may be damaged or encrypted.") from exc
    if not text.strip():
        raise JDImportError(
            "No selectable text found in this PDF (it may be a scanned image). Paste the JD text instead."
        )
    return text


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise JDImportError("The DOCX file could not be read. It may be damaged.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


# ---- URLs -------------------------------------------------------------------


def infer_platform(url: str) -> SourcePlatform:
    host = (urlsplit(url).hostname or "").lower()
    if host == "linkedin.com" or host.endswith(".linkedin.com"):
        return SourcePlatform.LINKEDIN
    if host == "naukri.com" or host.endswith(".naukri.com"):
        return SourcePlatform.NAUKRI
    return SourcePlatform.OTHER


def canonicalize_url(url: str) -> str:
    """Normalize a job URL for deduplication. Returns '' for blank or unparseable input.

    LinkedIn and Naukri postings reduce to their stable job ID; other URLs lose
    fragments, tracking parameters and trailing slashes.
    """
    url = (url or "").strip()
    if not url:
        return ""
    try:
        parts = urlsplit(url if "://" in url else f"https://{url}")
        host = (parts.hostname or "").lower()
        port = parts.port
    except ValueError:
        return ""
    if not host:
        return ""
    host = host.removeprefix("www.")
    path = parts.path or "/"
    query = parse_qsl(parts.query, keep_blank_values=False)

    platform = infer_platform(f"https://{host}")
    if platform is SourcePlatform.LINKEDIN:
        m = _LINKEDIN_JOB_RE.search(path)
        job_id = m.group(1) if m else dict(query).get("currentJobId", "")
        if job_id.isdigit():
            return f"https://linkedin.com/jobs/view/{job_id}"
    if platform is SourcePlatform.NAUKRI:
        m = _NAUKRI_JOB_RE.search(path.rstrip("/"))
        if m:
            return f"https://naukri.com/job-listings-{m.group(1)}"

    kept = sorted(
        (k, v) for k, v in query if k.lower() not in _TRACKING_PARAMS and not k.lower().startswith("utm_")
    )
    netloc = host if port in (None, 80, 443) else f"{host}:{port}"
    scheme = "https" if parts.scheme.lower() in ("http", "https") else parts.scheme.lower()
    return urlunsplit((scheme, netloc, path.rstrip("/") or "/", urlencode(kept), ""))


# ---- deduplication ----------------------------------------------------------


def _norm_field(value: str) -> str:
    return re.sub(r"[^\w]+", " ", value.casefold()).strip()


def description_hash(text: str) -> str:
    """Hash of the description with case, punctuation and whitespace ignored."""
    return hashlib.sha256(_norm_field(text).encode("utf-8")).hexdigest()


def _shingles(text: str, n: int = 3) -> set[tuple[str, ...]]:
    words = _norm_field(text).split()
    return {tuple(words[i:i + n]) for i in range(max(len(words) - n + 1, 1))} if words else set()


def text_similarity(a: str, b: str) -> float:
    """Jaccard similarity of 3-word shingles (0..1)."""
    sa, sb = _shingles(a), _shingles(b)
    if not sa or not sb:
        return 0.0
    return len(sa & sb) / len(sa | sb)


@dataclass(frozen=True)
class DuplicateMatch:
    job: Job
    reason: str


def find_duplicates(candidate: Job, existing: list[Job]) -> list[DuplicateMatch]:
    """Return stored jobs that are likely the same posting as ``candidate``.

    Cautious rules (any one suffices): same canonical URL; identical description
    text; or same company + title + location with near-identical descriptions.
    Different URLs alone never prove different postings, so text rules still apply.
    """
    matches: list[DuplicateMatch] = []
    cand_hash = description_hash(candidate.description)
    cand_key = tuple(_norm_field(v) for v in (candidate.company, candidate.title, candidate.location))
    for job in existing:
        if job.id == candidate.id:
            continue
        if candidate.canonical_url and job.canonical_url == candidate.canonical_url:
            matches.append(DuplicateMatch(job, "same job URL"))
        elif description_hash(job.description) == cand_hash:
            matches.append(DuplicateMatch(job, "identical description text"))
        elif all(cand_key[:2]) and cand_key == tuple(
            _norm_field(v) for v in (job.company, job.title, job.location)
        ):
            sim = text_similarity(candidate.description, job.description)
            if sim >= NEAR_DUPLICATE_SIMILARITY:
                matches.append(DuplicateMatch(job, f"same company/title/location and {sim:.0%} similar text"))
    return matches


def build_job(
    description: str,
    *,
    title: str = "",
    company: str = "",
    location: str = "",
    source_url: str = "",
    source_platform: Optional[SourcePlatform] = None,
    import_method=None,
    source_label: str = "",
    notes: Optional[list[str]] = None,
) -> Job:
    """Validate user-reviewed import fields into a ``Job`` (not yet saved)."""
    text, text_notes = check_jd_text(description)
    source_url = source_url.strip()
    kwargs = {}
    if import_method is not None:
        kwargs["import_method"] = import_method
    return Job(
        title=title,
        company=company,
        location=location,
        description=text,
        source_url=source_url,
        canonical_url=canonicalize_url(source_url),
        source_platform=source_platform or (infer_platform(source_url) if source_url else SourcePlatform.OTHER),
        source_label=source_label,
        import_notes=list(dict.fromkeys([*(notes or []), *text_notes])),
        **kwargs,
    )


def merge_import_into(existing: Job, incoming: Job) -> Job:
    """Update ``existing`` with a re-import of the same posting.

    The new JD text and source replace the old; user-owned fields (status, notes)
    are kept, and blank incoming metadata never erases what the user entered.
    """
    return existing.model_copy(update={
        "description": incoming.description,
        "source_url": incoming.source_url or existing.source_url,
        "canonical_url": incoming.canonical_url or existing.canonical_url,
        "source_platform": incoming.source_platform if incoming.source_url else existing.source_platform,
        "import_method": incoming.import_method,
        "source_label": incoming.source_label or existing.source_label,
        "title": incoming.title or existing.title,
        "company": incoming.company or existing.company,
        "location": incoming.location or existing.location,
        "import_notes": incoming.import_notes,
    })
