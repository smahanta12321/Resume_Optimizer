"""DOCX/PDF export and validation tests. Renders real files into tmp dirs. Fictional data only."""

from __future__ import annotations

import pytest

from src.schemas.export import CheckStatus
from src.schemas.job import Job
from src.schemas.profile import Bullet, Contact, Experience, MasterProfile, TargetLength, utc_now
from src.schemas.resume import ResumeContent, ResumeEntry, ResumeBullet
from src.services.export_service import ExportError, delete_job_outputs, export_version, safe_filename
from src.services.export_validation import normalize, validate_docx, validate_pdf
from src.services.jd_parser import parse_jd
from src.services.matching_service import build_match_report
from src.services.render.blocks import build_blocks, date_range, format_month
from src.services.render.docx_renderer import render_docx
from src.services.render.pdf_renderer import render_pdf
from src.services.render.style import CLASSIC, PdfFont, pdf_font, unrenderable_chars
from src.services.resume_generator import build_draft, new_version


@pytest.fixture
def approved(sample_profile):
    analysis = parse_jd(Job(description="Requirements\n• Python, SQL and Docker\n")).analysis
    report, _ = build_match_report(analysis, sample_profile)
    content, omitted = build_draft(sample_profile, analysis, report)
    v = new_version("job-abc123", sample_profile, report, content, omitted, generator="deterministic", role=None,
                    target_pages=TargetLength.ONE_PAGE)
    return v.model_copy(update={"approved_at": utc_now()})


def _status(record, name, file=None):
    return next(c.status for c in record.checks if c.name == name and (file is None or c.file == file))


# ---- blocks ------------------------------------------------------------------------


def test_dates_and_blocks(approved):
    assert format_month("2022-07") == "Jul 2022" and format_month(None) == ""
    assert date_range("2022-07", None) == "Jul 2022 – Present"
    assert date_range("2016-08", "2020-05", open_ended=False) == "Aug 2016 – May 2020"
    assert date_range("2016-08", None, open_ended=False) == "Aug 2016"
    blocks = build_blocks(approved.content)
    kinds = [b.kind for b in blocks]
    assert kinds[0] == "name" and "contact" in kinds
    headings = [b.text for b in blocks if b.kind == "heading"]
    assert headings == ["Summary", "Skills", "Experience", "Projects", "Education"]
    assert sum(k == "bullet" for k in kinds) == sum(len(e.bullets) for e in
                                                     [*approved.content.experience, *approved.content.projects])


# ---- full export -----------------------------------------------------------------


def test_export_approved_version_passes_all_checks(approved, tmp_path):
    record = export_version(approved, output_dir=tmp_path, company="Fictional Corp", role="Data Engineer")
    assert record.failed == [], [c.detail for c in record.failed]
    assert record.pdf_pages == 1 and record.target_pages == 1
    for path in (record.docx_path, record.pdf_path):
        p = tmp_path / "job-abc123" / approved.id
        assert str(p) in path
        assert path.endswith(("Fictional_Corp_Data_Engineer_Resume.docx", "Fictional_Corp_Data_Engineer_Resume.pdf"))
    assert _status(record, "DOCX and PDF match") is CheckStatus.PASS
    assert _status(record, "No tables") is CheckStatus.PASS
    assert _status(record, "Text extracts cleanly") is CheckStatus.PASS


def test_unapproved_version_cannot_export(approved, tmp_path):
    with pytest.raises(ExportError, match="Approve"):
        export_version(approved.model_copy(update={"approved_at": None}), output_dir=tmp_path)
    with pytest.raises(ExportError, match="page size"):
        export_version(approved, output_dir=tmp_path, page_size="Tabloid")


def test_letter_page_size(approved, tmp_path):
    import pdfplumber

    record = export_version(approved, output_dir=tmp_path, page_size="Letter")
    with pdfplumber.open(record.pdf_path) as pdf:
        assert round(pdf.pages[0].width) == 612  # 8.5in * 72


def test_overflow_is_flagged_not_shrunk(sample_profile, tmp_path):
    long = [ResumeBullet(evidence_ids=["bul-sample01"], text=f"Fictional achievement number {i} described in "
                         "enough words to wrap across the full width of the page for testing.", source_text="x")
            for i in range(60)]
    content = ResumeContent(contact=Contact(full_name="Fictional Person"), experience=[
        ResumeEntry(evidence_id="exp-sample01", kind="experience", title="Engineer", bullets=long)])
    blocks = build_blocks(content)
    path = render_pdf(blocks, tmp_path / "long.pdf", CLASSIC)
    checks, pages = validate_pdf(path, blocks, 1, pdf_font())
    assert pages >= 2
    fit = next(c for c in checks if c.name == "Fits target length")
    assert fit.status is CheckStatus.WARN and "never shrunk" in fit.detail
    assert next(c for c in checks if c.name == "Readable font size").status is CheckStatus.PASS


# ---- validators catch problems ----------------------------------------------------


def test_pdf_missing_content_detected(approved, tmp_path):
    blocks = build_blocks(approved.content)
    path = render_pdf(blocks[:-2], tmp_path / "partial.pdf", CLASSIC)  # drop the last lines
    checks, _ = validate_pdf(path, blocks, 1, pdf_font())
    assert next(c for c in checks if c.name == "All approved content present").status is CheckStatus.FAIL


def test_corrupt_files_fail_to_open(approved, tmp_path):
    blocks = build_blocks(approved.content)
    (tmp_path / "bad.pdf").write_bytes(b"%PDF-1.4 garbage")
    (tmp_path / "bad.docx").write_bytes(b"PK\x03\x04garbage")
    checks, pages = validate_pdf(tmp_path / "bad.pdf", blocks, 1, pdf_font())
    assert pages is None and checks[0].status is CheckStatus.FAIL
    assert validate_docx(tmp_path / "bad.docx", blocks)[0].status is CheckStatus.FAIL


def test_docx_table_and_missing_text_detected(approved, tmp_path):
    import docx

    blocks = build_blocks(approved.content)
    path = render_docx(blocks, tmp_path / "r.docx", CLASSIC)
    d = docx.Document(str(path))
    d.add_table(rows=1, cols=2)
    d.paragraphs[-1].text = "tampered"
    d.save(str(path))
    checks = {c.name: c.status for c in validate_docx(path, blocks)}
    assert checks["No tables"] is CheckStatus.FAIL
    assert checks["All approved content present"] is CheckStatus.FAIL


def test_docx_header_text_warns(approved, tmp_path):
    import docx

    blocks = build_blocks(approved.content)
    path = render_docx(blocks, tmp_path / "r.docx", CLASSIC)
    d = docx.Document(str(path))
    d.sections[0].header.paragraphs[0].text = "email@example.com"
    d.save(str(path))
    checks = {c.name: c.status for c in validate_docx(path, blocks)}
    assert checks["Contact details in body"] is CheckStatus.WARN


def test_fallback_font_flags_unsupported_characters():
    helv = PdfFont("Helvetica", "Helvetica-Bold", False, "built-in Helvetica")
    assert unrenderable_chars("Café – “quoted” • ok", helv) == set()
    assert unrenderable_chars("Résumé → नमस्ते", helv) >= {"→", "न"}
    assert unrenderable_chars("→", PdfFont("X", "X-Bold", True)) == set()


def test_normalize():
    assert normalize("Jul 2022 – Present\n• Built  “it”") == normalize("jul 2022 - present built \"it\"")


# ---- file safety --------------------------------------------------------------------


def test_safe_filename():
    assert safe_filename("Asha Example (FICTIONAL)", "Fictional Corp / Ltd.", "ML Engineer") == \
        "Asha_Example_FICTIONAL_Fictional_Corp_Ltd_ML_Engineer"
    assert safe_filename("../../etc", "") == "etc"
    assert safe_filename("नमस्ते") == "resume"


def test_path_traversal_refused(approved, tmp_path):
    with pytest.raises(ExportError):
        export_version(approved.model_copy(update={"job_id": "../evil"}), output_dir=tmp_path)


def test_delete_job_outputs(approved, tmp_path):
    export_version(approved, output_dir=tmp_path)
    assert (tmp_path / "job-abc123").is_dir()
    assert delete_job_outputs(tmp_path, "job-abc123") is True
    assert not (tmp_path / "job-abc123").exists()
    assert delete_job_outputs(tmp_path, "job-abc123") is False
    assert delete_job_outputs(tmp_path, "../x") is False


# ---- store -------------------------------------------------------------------------


def test_export_record_storage(tmp_path, approved):
    from src.persistence.job_store import JobStore, JobStoreError

    store = JobStore(tmp_path)
    job = store.add(Job(description="Fictional JD text"))
    v = store.add_resume_version(approved.model_copy(update={"job_id": job.id}))
    record = export_version(v.model_copy(update={"approved_at": utc_now()}), output_dir=tmp_path / "out")
    with pytest.raises(JobStoreError, match="approved"):
        store.save_export(record)
    store.approve_resume_version(v.id)
    store.save_export(record)
    assert store.get_export(v.id) == record
    assert store.export_counts() == {job.id: 1}
    store.delete(job.id)
    assert store.get_export(v.id) is None


def test_compact_template_exports_cleanly(approved, tmp_path):
    from src.services.render.style import COMPACT, MIN_FONT_PT

    assert min(COMPACT.body_pt, COMPACT.contact_pt, COMPACT.heading_pt) >= MIN_FONT_PT
    record = export_version(approved, output_dir=tmp_path, style=COMPACT)
    assert record.renderer == "compact-v1" and record.failed == []
