"""JD import tests: text cleanup, file extraction, URL canonicalization, dedup. Fictional data only."""

from __future__ import annotations

import io

import pytest

from src.schemas.job import MAX_JD_CHARS, ImportMethod, SourcePlatform
from src.services.jd_import_service import (
    MAX_UPLOAD_BYTES,
    JDImportError,
    build_job,
    canonicalize_url,
    check_jd_text,
    extract_text_from_file,
    find_duplicates,
    infer_platform,
    normalize_jd_text,
)

JD = (
    "Fictional Corp is hiring a Machine Learning Engineer to build retrieval systems. "
    "You will design data pipelines in Python, deploy models with Docker and Kubernetes, "
    "and work with product teams. Requirements: 3+ years of Python, experience with PyTorch, "
    "SQL, and cloud platforms. Nice to have: LLM fine-tuning experience."
)


# ---- text -------------------------------------------------------------------


def test_normalize_strips_controls_and_collapses_whitespace():
    raw = "Title\x00\r\n\r\n\r\n\r\nLine   with\t\tspaces here  \n\x07End"
    assert normalize_jd_text(raw) == "Title\n\nLine with spaces here\nEnd"


def test_check_jd_text_empty_long_and_short():
    with pytest.raises(JDImportError, match="empty"):
        check_jd_text(" \n\t ")
    with pytest.raises(JDImportError, match="too long"):
        check_jd_text("a" * (MAX_JD_CHARS + 1))
    text, notes = check_jd_text("Short JD")
    assert text == "Short JD" and notes and "short" in notes[0]
    _, notes = check_jd_text(JD)
    assert notes == []


# ---- files ------------------------------------------------------------------


def test_txt_extraction_utf8_and_cp1252():
    assert extract_text_from_file("jd.txt", "Café — role".encode("utf-8")) == "Café — role"
    assert extract_text_from_file("jd.TXT", "Café".encode("cp1252")) == "Café"


@pytest.mark.parametrize(
    "name,data,match",
    [
        ("jd.exe", b"MZ", "Unsupported"),
        ("jd", b"text", "Unsupported"),
        ("jd.txt", b"", "empty"),
        ("jd.txt", b"a" * (MAX_UPLOAD_BYTES + 1), "too large"),
        ("jd.txt", b"ab\x00cd", "binary"),
        ("jd.pdf", b"not a pdf", "does not look like a PDF"),
        ("jd.docx", b"not a zip", "does not look like a DOCX"),
        ("jd.pdf", b"%PDF-1.4 garbage", "could not be read"),
        ("jd.docx", b"PK\x03\x04garbage", "could not be read"),
    ],
    ids=["exe", "no-ext", "empty", "too-large", "binary-txt", "fake-pdf", "fake-docx", "bad-pdf", "bad-docx"],
)
def test_file_rejections(name, data, match):
    with pytest.raises(JDImportError, match=match):
        extract_text_from_file(name, data)


def test_docx_extraction_includes_tables():
    import docx

    d = docx.Document()
    d.add_paragraph("Fictional ML Engineer role")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Location"
    table.rows[0].cells[1].text = "Remote"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text_from_file("jd.docx", buf.getvalue())
    assert "Fictional ML Engineer role" in text
    assert "Location | Remote" in text


def _make_pdf(text: str | None) -> bytes:
    """Build a tiny one-page PDF by hand (no extra dependency). ``None`` -> page with no text."""
    content = b"" if text is None else f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, obj in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def test_pdf_extraction():
    assert "Fictional Data Scientist" in extract_text_from_file("jd.pdf", _make_pdf("Fictional Data Scientist"))


def test_pdf_without_text_suggests_paste():
    with pytest.raises(JDImportError, match="scanned"):
        extract_text_from_file("jd.pdf", _make_pdf(None))


# ---- URLs -------------------------------------------------------------------


@pytest.mark.parametrize(
    "url,expected",
    [
        ("https://www.linkedin.com/jobs/view/3987654321/?trk=abc&refId=x", "https://linkedin.com/jobs/view/3987654321"),
        ("https://in.linkedin.com/jobs/view/ml-engineer-at-fictional-corp-3987654321",
         "https://linkedin.com/jobs/view/3987654321"),
        ("https://www.linkedin.com/jobs/search/?currentJobId=3987654321&keywords=ml",
         "https://linkedin.com/jobs/view/3987654321"),
        ("https://www.naukri.com/job-listings-data-scientist-fictional-corp-bengaluru-3-to-5-years-120924500123?src=jobsearchDesk",
         "https://naukri.com/job-listings-120924500123"),
        ("HTTP://Careers.Example.com/jobs/42/?utm_source=x&b=2&a=1#apply", "https://careers.example.com/jobs/42?a=1&b=2"),
        ("careers.example.com/jobs/42", "https://careers.example.com/jobs/42"),
        ("https://example.com:8443/x", "https://example.com:8443/x"),
        ("", ""),
        ("   ", ""),
        ("https://[bad", ""),
    ],
)
def test_canonicalize_url(url, expected):
    assert canonicalize_url(url) == expected


def test_infer_platform():
    assert infer_platform("https://in.linkedin.com/jobs/view/1") is SourcePlatform.LINKEDIN
    assert infer_platform("https://www.naukri.com/x") is SourcePlatform.NAUKRI
    assert infer_platform("https://notlinkedin.com/x") is SourcePlatform.OTHER


# ---- build + dedup ------------------------------------------------------------


def test_build_job_sets_derived_fields():
    job = build_job(JD, title="ML Engineer", source_url="https://www.linkedin.com/jobs/view/3987654321/?trk=x",
                    import_method=ImportMethod.URL, notes=["from url"])
    assert job.source_platform is SourcePlatform.LINKEDIN
    assert job.canonical_url == "https://linkedin.com/jobs/view/3987654321"
    assert job.import_method is ImportMethod.URL
    assert job.import_notes == ["from url"]


def test_build_job_rejects_empty():
    with pytest.raises(JDImportError):
        build_job("  ")


def test_duplicate_by_canonical_url():
    a = build_job(JD, source_url="https://www.linkedin.com/jobs/view/3987654321/")
    b = build_job("Completely different text " * 20, source_url="https://linkedin.com/jobs/view/3987654321?trk=1")
    [m] = find_duplicates(b, [a])
    assert m.job.id == a.id and m.reason == "same job URL"


def test_duplicate_by_identical_text_ignoring_case_and_punctuation():
    a = build_job(JD)
    b = build_job(JD.upper().replace(".", " "))
    [m] = find_duplicates(b, [a])
    assert m.reason == "identical description text"


def test_near_duplicate_requires_matching_company_and_title():
    a = build_job(JD, company="Fictional Corp", title="ML Engineer", location="Remote")
    edited = JD + " Apply by Friday."
    same = build_job(edited, company="fictional corp", title="ML  Engineer", location="remote")
    assert [m.job.id for m in find_duplicates(same, [a])] == [a.id]
    other_company = build_job(edited, company="Other Fictional Inc", title="ML Engineer", location="Remote")
    assert find_duplicates(other_company, [a]) == []
    no_metadata = build_job(edited)
    assert find_duplicates(no_metadata, [build_job(JD)]) == []


def test_unrelated_jobs_not_duplicates():
    a = build_job(JD, company="Fictional Corp", title="ML Engineer")
    b = build_job("We need a pastry chef for a fictional bakery. " * 10, company="Fictional Corp", title="ML Engineer")
    assert find_duplicates(b, [a]) == []


def test_job_is_not_its_own_duplicate():
    a = build_job(JD)
    assert find_duplicates(a, [a]) == []


def test_merge_import_keeps_user_fields_and_fills_blanks():
    from src.schemas.job import JobStatus
    from src.services.jd_import_service import merge_import_into

    existing = build_job(JD, title="ML Engineer", company="Fictional Corp").model_copy(
        update={"status": JobStatus.APPLIED, "notes": "my note"})
    incoming = build_job(JD + " Updated.", source_url="https://www.naukri.com/job-listings-x-120924500123",
                         import_method=ImportMethod.URL)
    merged = merge_import_into(existing, incoming)
    assert merged.id == existing.id
    assert merged.description.endswith("Updated.")
    assert (merged.title, merged.company) == ("ML Engineer", "Fictional Corp")
    assert (merged.status, merged.notes) == (JobStatus.APPLIED, "my note")
    assert merged.source_platform is SourcePlatform.NAUKRI and merged.import_method is ImportMethod.URL
